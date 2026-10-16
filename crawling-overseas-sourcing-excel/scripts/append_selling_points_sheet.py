import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
from urllib.parse import urlparse
import json
import re
from difflib import SequenceMatcher

EXCEL_PATH = Path(r"C:/Users/imda0/Desktop/해외구매대행_사입_TOP50_상품리뷰_수집분석.xlsx")
JSON_PATH = Path(r"C:/tmp/product_selling_points.json")
SOURCE_EXCEL_PATH = Path(r"C:/Users/imda0/Desktop/해외구매대행_사입_추천상품_TOP50.xlsx")

def _normalize_name(value):
    return "".join(ch.casefold() for ch in str(value or "") if ch.isalnum())


def _direct_1688_offer(url):
    """Accept only a direct 1688 offer detail page, never a keyword search URL."""
    parsed = urlparse(str(url or "").strip())
    if parsed.hostname != "detail.1688.com":
        return ""
    return str(url).strip() if re.fullmatch(r"/offer/\d+\.html", parsed.path) else ""


def _load_product_links():
    """Load source URLs, rejecting 1688 search pages as product links."""
    if not SOURCE_EXCEL_PATH.is_file():
        raise FileNotFoundError(f"원본 추천 엑셀을 찾을 수 없습니다: {SOURCE_EXCEL_PATH}")
    source_wb = openpyxl.load_workbook(SOURCE_EXCEL_PATH, data_only=True, read_only=False)
    if "전체 통합 TOP 50" not in source_wb.sheetnames:
        raise KeyError("원본 엑셀에 '전체 통합 TOP 50' 시트가 없습니다.")
    source_ws = source_wb["전체 통합 TOP 50"]
    links = []
    for row in range(2, source_ws.max_row + 1):
        name = str(source_ws.cell(row, 5).value or "").strip()
        if not name:
            continue
        def value_or_hyperlink(col):
            cell = source_ws.cell(row, col)
            return (cell.hyperlink.target if cell.hyperlink else cell.value) or ""
        links.append({
            "name": name,
            "coupang": str(value_or_hyperlink(18)).strip(),
            "1688": _direct_1688_offer(value_or_hyperlink(20)),
            "naver": str(value_or_hyperlink(21)).strip(),
        })
    return links


def _match_product_links(product_name, link_rows):
    target = _normalize_name(product_name)
    if not target:
        return {"coupang": "", "1688": "", "naver": "", "matched_name": ""}
    ranked = []
    for item in link_rows:
        candidate = _normalize_name(item["name"])
        if not candidate:
            continue
        if target == candidate:
            score = 2.0
        elif target in candidate or candidate in target:
            score = 1.0 + min(len(target), len(candidate)) / max(len(target), len(candidate))
        else:
            score = SequenceMatcher(None, target, candidate).ratio()
        ranked.append((score, item))
    if not ranked:
        return {"coupang": "", "1688": "", "naver": "", "matched_name": ""}
    score, best = max(ranked, key=lambda pair: pair[0])
    if score < 0.72:
        return {"coupang": "", "1688": "", "naver": "", "matched_name": ""}
    return {**best, "matched_name": best["name"], "match_score": score}


def append_selling_points_sheet():
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)
    products = data.get("products", [])
    product_links = _load_product_links()

    wb = openpyxl.load_workbook(EXCEL_PATH)
    sheet_title = "상품별 핵심 소구점 및 카피"
    if sheet_title in wb.sheetnames:
        del wb[sheet_title]
    
    ws = wb.create_sheet(title=sheet_title)

    # Styles
    title_font = Font(name="맑은 고딕", size=14, bold=True, color="1F4E79")
    header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    body_font = Font(name="맑은 고딕", size=9)
    bold_body_font = Font(name="맑은 고딕", size=9, bold=True)
    hook_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")
    border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = [
        "순번", "추천 순위", "상품명", "카테고리",
        "🛒 쿠팡 상품 URL (클릭)", "🇨🇳 1688 오퍼 URL (클릭)", "🔍 네이버 가격비교 URL (클릭)",
        "근거 신뢰도",
        "킬러 메인 헤드라인 (클릭 유발)", "서브 설명 카피 (실질 혜택)",
        "📋 원클릭 마케팅 기획 종합본 (복사용)",
        "고객 핵심 결핍 (Pain Point)", "제품 핵심 해결책 (Solution)",
        "3대 핵심 소구점 (Appeal Points)", "구매 전환 트리거 (결제 결심 이유)",
        "실구매 고객 리얼 인용구", "표시광고 주의사항 (과장 방지)"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    cat_map = {
        "MERAKA만능 채칼": "주방용품",
        "동물 자석 커플 키링": "팬시/잡화",
        "27리빙 여행용 깔끔한 소분": "여행/생활",
        "비비돔 무선 미니 고데기": "디지털/이미용",
        "의자 커버": "홈/인테리어",
        "파스텔펫 강아지": "반려동물",
        "Zeekr 차량용": "차량용품",
        "베베데코 고탄력": "홈/인테리어"
    }

    for idx, p in enumerate(products, start=1):
        p_name = p.get("product", "")
        cat_val = "생활잡화"
        for k, v in cat_map.items():
            if k in p_name:
                cat_val = v
                break

        links = _match_product_links(p_name, product_links)
        if not links.get("matched_name"):
            print(f"[WARN] 추천 원본 링크 매칭 실패: {p_name}")

        # Refined headlines by Antigravity recommendations
        hl = p.get("headline", "")
        if "손 다칠 걱정 끝" in hl: hl = "푸드홀더로 안전하게, 균일한 채썰기"
        elif "뜬 앞머리 10초 만에" in hl: hl = "바쁜 아침 뜬 앞머리 간편 정돈"
        elif "여행 짐 부피 반으로" in hl: hl = "짐 부피를 쏙 줄이는 소분 파우치"
        elif "소파 새로 사지 말고" in hl: hl = "소파에 씌워 새것처럼 분위기 전환"

        appeal_text = "\n".join(f"{i+1}. {pt}" for i, pt in enumerate(p.get("appeal_points", [])))
        quotes_text = "\n".join(f'"{q}"' for q in p.get("evidence_quotes", []))
        avoid_text = "\n".join(f"• {a}" for a in p.get("avoid_claims", []))

        # 원클릭 통합 복사본 텍스트 조립
        all_in_one = [
            f"【 {p_name} - 마케팅 기획 요약 】",
            f"■ 킬러 헤드라인: {hl}",
            f"■ 서브 카피: {p.get('sub_copy', '')}",
            f"■ 원본 상품/쿠팡 구매 링크: {links.get('coupang') or '원본 링크 미확인'}",
            f"■ 1688 실제 사입 상품 상세페이지: {links.get('1688') or '정확 일치 오퍼 미확인'}",
            f"■ 네이버 가격비교: {links.get('naver') or '비교 링크 미확인'}",
            "",
            "■ 고객 핵심 결핍 (Pain Point):",
            f"  {p.get('pain_point', '')}",
            "",
            "■ 제품 핵심 해결책 (Solution):",
            f"  {p.get('solution', '')}",
            "",
            "■ 3대 핵심 소구점 (Selling Points):",
            appeal_text,
            "",
            "■ 구매 전환 트리거 (결제 결심 이유):",
            f"  {p.get('purchase_trigger', '')}",
            "",
            "■ 실구매 고객 리얼 후기 인용:",
            quotes_text,
            "",
            "■ 과장 방지 및 상세페이지 주의사항:",
            avoid_text
        ]
        all_in_one_text = "\n".join(all_in_one)

        row_vals = [
            idx,
            f"{idx}위",
            p_name,
            cat_val,
            links.get("coupang") or "원본 링크 미확인",
            links.get("1688") or "직접 상품 확인 안 됨",
            links.get("naver") or "네이버 링크 미확인",
            p.get("evidence_confidence", "보통"),
            hl,
            p.get("sub_copy", ""),
            all_in_one_text,
            p.get("pain_point", ""),
            p.get("solution", ""),
            appeal_text,
            p.get("purchase_trigger", ""),
            quotes_text,
            avoid_text
        ]
        ws.append(row_vals)
        r_num = idx + 1
        ws.row_dimensions[r_num].height = 220

        for c_idx in range(1, len(row_vals) + 1):
            cell = ws.cell(r_num, c_idx)
            cell.border = border
            if c_idx in (5, 6, 7):
                link_value = row_vals[c_idx - 1]
                link_url = links.get({5: "coupang", 6: "1688", 7: "naver"}[c_idx], "")
                if link_url:
                    cell.hyperlink = link_url
                    cell.font = Font(name="맑은 고딕", size=9, color="0563C1", underline="single")
                else:
                    cell.font = body_font
            elif c_idx == 9:
                cell.font = hook_font
            elif c_idx in (2, 8):
                cell.font = bold_body_font
            else:
                cell.font = body_font

            align_h = "center" if c_idx in (1, 2, 4, 8) else "left"
            cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=True)
            if c_idx == 11:
                cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # 연노랑 복사강조

    # Column widths
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 11
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 42
    ws.column_dimensions["F"].width = 48
    ws.column_dimensions["G"].width = 48
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 32
    ws.column_dimensions["J"].width = 34
    ws.column_dimensions["K"].width = 55  # 📋 원클릭 복사본 열 넓게
    ws.column_dimensions["L"].width = 38
    ws.column_dimensions["M"].width = 38
    ws.column_dimensions["N"].width = 45
    ws.column_dimensions["O"].width = 38
    ws.column_dimensions["P"].width = 40
    ws.column_dimensions["Q"].width = 35

    matched = sum(1 for p in products if _match_product_links(p.get("product", ""), product_links).get("matched_name"))
    print(f"[LINKS] 원본 상품 링크 매칭: {matched}/{len(products)}개")

    try:
        wb.save(EXCEL_PATH)
        print(f"SUCCESS: Saved sheet to {EXCEL_PATH}")
    except PermissionError:
        import time
        alt_path = EXCEL_PATH.with_name(f"{EXCEL_PATH.stem}_{time.strftime('%H%M%S')}{EXCEL_PATH.suffix}")
        wb.save(alt_path)
        print(f"SUCCESS: Saved sheet to alt path {alt_path}")

if __name__ == "__main__":
    append_selling_points_sheet()
