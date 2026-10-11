import sys
import io
import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def build_review_workbook(data, out_path):
    wb = openpyxl.Workbook()
    # Sheet 1: All reviews
    ws1 = wb.active
    ws1.title = "상품별 실구매 리뷰 전수"

    header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    body_font = Font(name="맑은 고딕", size=9)
    star_font = Font(name="맑은 고딕", size=10, bold=True, color="ED7D31")
    border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers1 = [
        "순번", "추천 순위", "상품명", "카테고리", "플랫폼",
        "리뷰 번호", "작성자", "별점", "작성일자", "리뷰 원문 내용", "핵심 만족 포인트", "개선/주의 사항"
    ]
    ws1.append(headers1)
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_count = 1
    total_reviews = 0
    insights = []

    for item in data:
        p_name = item.get("name", "")
        p_cat = item.get("category", "")
        p_rank = item.get("rank", 0)
        p_plat = item.get("platform", "쿠팡")
        reviews = item.get("reviews", [])
        
        # Calculate insights
        good_points, bad_points = [], []

        for idx, rev in enumerate(reviews):
            row_count += 1
            total_reviews += 1
            author = rev.get("author", "구매자")
            date = rev.get("date", "2026.09")
            rating = rev.get("rating", 5)
            content = rev.get("content", "")

            # Simple sentiment tagging
            good = "만족 / 실용성 우수"
            bad = "없음 / 정상 사용"
            if "작" in content or "크" in content or "사이즈" in content:
                bad = "사이즈/용량 확인 필요"
            if "안전" in content or "손" in content:
                good = "안전성 / 손보호 만족"
            if "예쁘" in content or "귀엽" in content:
                good = "디자인 / 비주얼 만족"
            if "편" in content or "쉽" in content or "빠르" in content:
                good = "편리함 / 시간 절약"

            if good not in good_points: good_points.append(good)
            if bad != "없음 / 정상 사용" and bad not in bad_points: bad_points.append(bad)

            row_vals = [
                total_reviews, p_rank, p_name, p_cat, p_plat,
                idx + 1, author, f"★ {rating}.0", date, content, good, bad
            ]
            ws1.append(row_vals)
            for c_idx in range(1, len(row_vals) + 1):
                c = ws1.cell(row_count, c_idx)
                c.font = star_font if c_idx == 8 else body_font
                c.border = border
                c.alignment = Alignment(
                    horizontal="center" if c_idx in (1, 2, 4, 5, 6, 7, 8, 9) else "left",
                    vertical="center",
                    wrap_text=True
                )

        insights.append({
            "rank": p_rank,
            "name": p_name,
            "category": p_cat,
            "platform": p_plat,
            "review_count": len(reviews),
            "good_points": " · ".join(good_points) if good_points else "우수한 사용감과 높은 만족도",
            "bad_points": " · ".join(bad_points) if bad_points else "특이 불만 없음"
        })

    ws1.column_dimensions["A"].width = 7
    ws1.column_dimensions["B"].width = 10
    ws1.column_dimensions["C"].width = 30
    ws1.column_dimensions["D"].width = 12
    ws1.column_dimensions["E"].width = 10
    ws1.column_dimensions["F"].width = 10
    ws1.column_dimensions["G"].width = 12
    ws1.column_dimensions["H"].width = 10
    ws1.column_dimensions["I"].width = 13
    ws1.column_dimensions["J"].width = 65
    ws1.column_dimensions["K"].width = 25
    ws1.column_dimensions["L"].width = 25

    # Sheet 2: Insights
    ws2 = wb.create_sheet(title="상품별 VOC 인사이트 요약")
    headers2 = ["순위", "상품명", "카테고리", "플랫폼", "수집 리뷰 수", "고객 주요 만족 포인트", "주의 및 개선 피드백", "상세페이지 공략 포인트"]
    ws2.append(headers2)
    header_fill2 = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill2
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, ins in enumerate(insights, start=2):
        row_vals = [
            ins["rank"], ins["name"], ins["category"], ins["platform"],
            f"{ins['review_count']}개 수집", ins["good_points"], ins["bad_points"],
            f"{ins['good_points']} 전면 부각 + {ins['bad_points']} 선제적 안내"
        ]
        ws2.append(row_vals)
        for c_idx in range(1, len(row_vals) + 1):
            c = ws2.cell(r_idx, c_idx)
            c.font = body_font
            c.border = border
            c.alignment = Alignment(horizontal="center" if c_idx in (1, 3, 4, 5) else "left", vertical="center", wrap_text=True)

    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 32
    ws2.column_dimensions["C"].width = 12
    ws2.column_dimensions["D"].width = 10
    ws2.column_dimensions["E"].width = 14
    ws2.column_dimensions["F"].width = 35
    ws2.column_dimensions["G"].width = 30
    ws2.column_dimensions["H"].width = 40

    out_file = Path(out_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out_file)
        saved = out_file
    except PermissionError:
        import time
        saved = out_file.with_name(f"{out_file.stem}_{time.strftime('%H%M%S')}{out_file.suffix}")
        wb.save(saved)
        print(f"[안내] 파일이 열려 있어 다음 경로로 저장했습니다: {saved}")

    print(f"[성공] 총 {len(data)}개 상품, {total_reviews}개 리뷰 수집 엑셀 생성 완료: {saved}")
    return str(saved)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python build_review_excel.py <input.json> <output.xlsx>")
        sys.exit(1)
    with open(sys.argv[1], "r", encoding="utf-8") as f:
        data = json.load(f)
    build_review_workbook(data, sys.argv[2])
