#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
v4.0 Executive 4-Sheet Excel Workbook Builder (excel_builder.py)
- Sheet 1: 구매대행 추천 (All ranked purchasing-agency items)
- Sheet 2: 사입 추천 (All ranked direct-sourcing items)
- Sheet 3: 전체 통합 TOP 50 (All ranked items sorted by s_v3; sheet name kept for contract compatibility)
- Sheet 4: 소싱 분석 요약 및 규제 가이드 (Executive summary & KIPRIS/KC/식약처/1688/마진/키워드 가이드)
- 33-column layout with KIPRIS trademark & private-label judgment, clean rebrand title, Coupang spec/material summary,
  low-cost KC/MFDS certification guide, and bilingual 1688 factory inquiry message.
"""
import io
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

HEADERS = [
    "순번",
    "상품 사진 (72x72)",
    "소싱등급",
    "소싱 구분",
    "선정 상품명 (핵심 아이템)",
    "카테고리",
    "v3 소싱점수 (100점)",
    "바이럴(25)",
    "참신성(25)",
    "링크검증(20)",
    "소싱타당성(30)",
    "국내 시세 / 권장 판매가",
    "목표 마진율 및 단가 가이드",
    "통관·KC·식약처 안전도",
    "추천 소싱채널 & 2단계 스케일업 전략",
    "핵심 추천 사유 & 셀링포인트",
    "추천 검색 키워드 (네이버·쿠팡)",
    "① 구매/상점 원본 링크 (클릭)",
    "② 알리익스프레스 도매 검색 (클릭)",
    "③ 1688 중국 사입처 검색 (클릭)",
    "④ 네이버 가격비교/시세 확인 (클릭)",
    "⑤ 대표 이미지 원본 URL (HTTP 200 클릭)",
    "⑥ 발굴 원본 게시물 링크 (클릭)",
    "링크 검증 상태",
    "출처/마켓",
    "소셜 반응 지표",
    "원문 요약 발췌",
    "키프리스(KIPRIS) 상표권 & 내 브랜드(택갈이) 판정",
    "내 브랜드 등록용 클린 상품명 (상표 제거)",
    "쿠팡·국내 실측 재질·성분표·규격 & 고시 요약",
    "저비용 KC·식약처·성분 인증 돌파 가이드",
    "1688 중국 공장 문의 메시지 (중문+한글 해석)",
    "⑦ 키프리스(KIPRIS) 상표권 검색 (클릭)",
]

COL_WIDTHS = [
    6, 12, 13, 15, 42, 12, 13, 10, 10, 11, 12, 26, 48, 36, 52, 48, 34, 42, 40, 38, 40, 46, 38, 28, 14, 24, 44,
    40, 40, 50, 56, 62, 38,
]

assert len(HEADERS) == len(COL_WIDTHS) == 33, f"Column mismatch: {len(HEADERS)} != {len(COL_WIDTHS)}"


def build_v3_workbook(top_agency: list[dict], top_sourcing: list[dict], out_path: Path, cache_dir: Path) -> dict:
    if not top_agency or not top_sourcing:
        raise ValueError(f"Expected non-empty verified rows, got {len(top_agency)}/{len(top_sourcing)}")
    names = [str(r.get("clean_name", "")).strip().casefold() for r in top_agency + top_sourcing]
    if any(not n for n in names) or len(set(names)) != len(names):
        raise ValueError("Workbook requires non-empty unique product names")
    for r in top_agency + top_sourcing:
        if not isinstance(r.get("thumb_png"), (bytes, bytearray)) or len(r["thumb_png"]) < 128:
            raise ValueError(f"Missing verified thumbnail for {r.get('clean_name', '<unknown>')}")
    all_50 = sorted(top_agency + top_sourcing, key=lambda x: x["s_v3"], reverse=True)
    wb = openpyxl.Workbook()
    default_ws = wb.active

    header_fill_agency = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_fill_sourcing = PatternFill(start_color="275317", end_color="275317", fill_type="solid")
    header_fill_all = PatternFill(start_color="262626", end_color="262626", fill_type="solid")
    fill_grade_s = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_grade_a = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_reg_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_reg_yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_reg_orange = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="맑은 고딕", size=9)
    bold_font = Font(name="맑은 고딕", size=9, bold=True)
    url_font = Font(name="맑은 고딕", size=8, color="0563C1", underline="single")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    def populate_sheet(ws, rows, header_fill):
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "F2"
        ws.append(HEADERS)
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for idx, r in enumerate(rows, 1):
            row_vals = [
                idx,
                "",
                r["grade"],
                r["sourcing_type"],
                r["clean_name"],
                r["category"],
                round(float(r["s_v3"]), 1),
                round(float(r["v_norm"]), 1),
                round(float(r["n_norm"]), 1),
                round(float(r["l_bonus"]), 1),
                round(float(r["f_sourcing"]), 1),
                r["price_str"],
                r["margin_guide"],
                r["reg_badge"],
                r["channel_strategy"],
                r.get("recommend_reason") or r.get("selling_point") or "",
                r.get("seo_keywords") or r.get("search_kw") or "",
                r["purchase_url"],
                r["ali_sourcing_url"],
                r["s1688_url"],
                r["naver_check_url"],
                r["primary_image_url"],
                r["post_url"],
                r["link_status"],
                r["merchant"],
                r["engagement"],
                r["raw_excerpt"],
                r.get("kipris_status", ""),
                r.get("rebrand_clean_name", ""),
                r.get("spec_material_summary", ""),
                r.get("cert_cost_saving_guide", ""),
                r.get("china_1688_message", ""),
                r.get("kipris_search_url", ""),
            ]
            ws.append(row_vals)
            row_num = idx + 1
            ws.row_dimensions[row_num].height = 58

            if r.get("thumb_png"):
                xl_img = XLImage(io.BytesIO(r["thumb_png"]))
                xl_img.width = 72
                xl_img.height = 72
                ws.add_image(xl_img, f"B{row_num}")

            for col_idx in range(1, len(HEADERS) + 1):
                c = ws.cell(row=row_num, column=col_idx)
                c.border = thin_border
                if col_idx in (7, 8, 9, 10, 11):
                    c.number_format = "0.0"
                if col_idx in (1, 2, 3, 4, 6, 7, 8, 9, 10, 11, 24, 25):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.font = bold_font if col_idx in (1, 3, 7) else body_font
                    if col_idx == 3:
                        if "S급" in str(c.value):
                            c.fill = fill_grade_s
                        elif "A급" in str(c.value):
                            c.fill = fill_grade_a
                elif col_idx == 14:
                    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    c.font = bold_font
                    val_s = str(c.value or "")
                    if "🟢" in val_s:
                        c.fill = fill_reg_green
                    elif "🟠" in val_s:
                        c.fill = fill_reg_orange
                    else:
                        c.fill = fill_reg_yellow
                elif col_idx == 28:
                    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    c.font = bold_font
                    val_s = str(c.value or "")
                    if "🟢" in val_s or "✅" in val_s:
                        c.fill = fill_reg_green
                    elif "🔵" in val_s:
                        c.fill = fill_reg_yellow
                    elif "⚠️" in val_s:
                        c.fill = fill_reg_orange
                elif col_idx in (18, 19, 20, 21, 22, 23, 33):
                    c.alignment = Alignment(horizontal="left", vertical="center", shrink_to_fit=False)
                    c.font = url_font
                    val_str = str(c.value or "").strip()
                    if val_str.startswith("http"):
                        c.hyperlink = val_str
                else:
                    wrap = col_idx in (5, 12, 13, 15, 16, 17, 27, 28, 29, 30, 31, 32)
                    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=wrap)
                    c.font = bold_font if col_idx in (5, 12, 16, 29) else body_font
        for col_idx, width in enumerate(COL_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"

    ws_agency = wb.create_sheet(title="구매대행 추천")
    populate_sheet(ws_agency, top_agency, header_fill_agency)

    ws_sourcing = wb.create_sheet(title="사입 추천")
    populate_sheet(ws_sourcing, top_sourcing, header_fill_sourcing)

    ws_all = wb.create_sheet(title="전체 통합 TOP 50")
    populate_sheet(ws_all, all_50, header_fill_all)

    ws_guide = wb.create_sheet(title="소싱 분석 요약 및 규제 가이드")
    ws_guide.views.sheetView[0].showGridLines = True
    guide_rows = [
        ["항목 / 구분", "핵심 기준 및 수치", "실전 운영 전략 및 주의사항"],
        [
            "1. v4.0 소싱점수(100점) 산출 공식",
            "S_v3 = (바이럴25 + 참신성25 + 링크검증20 + 소싱타당성30) × 규제안전계수",
            "인스타 댓글 낚시(댓글 > 좋아요×1.8) 감쇠 보정 적용 및 공식 상점/가격 확인·고수요 상품 가산점 부여 (S급 82점↑ / A급 74점↑ / B급 65점↑)",
        ],
        [
            "2. 해외구매대행 추천 (전수 순위) 기준",
            "전기/배터리/무선(KC·전파법) 또는 고단가(3.0만원↑) · 식품접촉 테스트 품목",
            "개인통관고유부호(PCC) 기반 1인 1대 전파법/KC 면제 활용 무재고 테스트 후 월 30건 이상 판매 시 정식수입 전환",
        ],
        [
            "3. 사입 추천 (전수 순위) 기준",
            "비전기·비전파·무인증 공산품 · 경량/소형 · 반복수요/묶음판매 고마진 품목",
            "1단계 타오바오/알리 10~20개 샘플 검수 → 2단계 1688 100~300개 LCL 사입 + 쿠팡 로켓그로스/스마트스토어 입점",
        ],
        [
            "4. 실전 단가·마진율 산출 기준",
            "쿠팡 실측가 + 네이버 쇼핑 실시간 시세 연동 / 1688 도매원가(¥) 역산",
            "구매대행 목표 마진율 25~38% (관부가세 $150 면세) / 사입 목표 마진율 40~58% (1688 원가 20%대 + 2~3개 세트 구성으로 객단가 극대화)",
        ],
        [
            "5. 핵심 추천 사유 & 검색 키워드 활용",
            "16열 [핵심 추천 사유 & 셀링포인트] / 17열 [추천 검색 키워드 (네이버·쿠팡)]",
            "상세페이지 최상단 후킹 카피 및 썸네일 문구로 16열 셀링포인트를 활용하고, 스마트스토어/쿠팡 상품명 및 검색어 태그에 17열 키워드 즉시 등록",
        ],
        [
            "6. 식약처(MFDS) 식품용 기구 주의사항",
            "프라이팬, 식기, 거름망, 텀블러 등 음식물에 직접 닿는 주방용품",
            "사입(정식수입) 시 재질·색상별 식약처 정밀검사비(30~80만원) 발생 → 반드시 구매대행으로 먼저 수요 검증 권장",
        ],
        [
            "7. 영구 이미지 썸네일 캐시 & 폴백 배지",
            str(cache_dir),
            "인스타/스레드 CDN 만료(403) 실시간 복구 + 로컬 PNG 캐시 및 외부 CDN 장애 시 72x72 규격 폴백 배지 자동 생성으로 전 후보 누락 방지",
        ],
        [
            "8. 키프리스(KIPRIS) 상표권 & 내 브랜드(택갈이) 전략",
            "28열 [상표권 & 택갈이 판정] · 29열 [클린 상품명] · 33열 [키프리스 검색 링크]",
            "🟢 무브랜드 공용상품 및 ✅ 국내 단순 리셀러 상표([리브나인], 투데이리빙 등)는 상표명을 제거한 29열 클린 상품명에 [내브랜드]를 붙여 즉시 독점 상품화 가능. 🔵 호환용 액세서리는 'OO 호환용' 표기 준수, ⚠️ 글로벌 원천 브랜드는 택갈이 금지(정품 구매대행 전용)",
        ],
        [
            "9. 쿠팡 실측(insane-search) 기반 저비용 인증 돌파 가이드",
            "30열 [쿠팡 실측 재질·성분·규격 요약] · 31열 [저비용 KC·식약처·성분 인증 가이드]",
            "쿠팡 상세페이지/고시정보의 기존 KC인증번호로 안전코리아·국립전파연구원에서 중국 원천 제조공장을 역추적해 동일·파생모델 인증비 절감. 식약처(impfood.mfds.go.kr) 기등록 해외제조업소 코드 활용 및 단일색상·100% 재질성분표 확보로 정밀검사비 최소화, 생활화학제품은 CAS 전성분표·MSDS 사전 확보",
        ],
        [
            "10. 1688 중국 공장 왕왕(WangWang)·위챗 문의 템플릿",
            "32열 [1688 중국 공장 문의 메시지 (중문 원문 + 한국어 해석)]",
            "상품별 카테고리·재질에 맞춘 중문 원문(복붙용)과 한국어 해석 동시 수록: ① 100% 재질성분표(材质成分表)/MSDS 요청 ② KC·CE·RoHS·FDA·FCC 성적서 보유 확인 ③ 공용금형(公模)/독점금형(私模) 지재권 확인 ④ 무지 패키지 및 OEM 로고 각인 MOQ 협상",
        ],
    ]
    for r_idx, gr in enumerate(guide_rows, 1):
        ws_guide.append(gr)
        ws_guide.row_dimensions[r_idx].height = 26
        for c_idx in range(1, 4):
            cell = ws_guide.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            if r_idx == 1:
                cell.fill = header_fill_all
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = bold_font if c_idx == 1 else body_font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws_guide.column_dimensions["A"].width = 34
    ws_guide.column_dimensions["B"].width = 52
    ws_guide.column_dimensions["C"].width = 84

    wb.remove(default_ws)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    verified_cnt = sum(1 for x in all_50 if x.get("img_verified_200"))
    return {
        "out_path": str(out_path),
        "file_size": out_path.stat().st_size,
        "agency_count": len(top_agency),
        "sourcing_count": len(top_sourcing),
        "total_count": len(all_50),
        "verified_images": verified_cnt,
    }
